import io
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from app.services.csv_import import (
    MAX_DATA_ROWS,
    MAX_FILE_SIZE_BYTES,
    ParsedImportData,
)

class XlsxImportError(ValueError):
    """Raised when an XLSX file cannot be safely read for import."""


def _load_workbook(file_bytes: bytes):
    if not file_bytes:
        raise XlsxImportError("The XLSX file is empty.")

    if len(file_bytes) > MAX_FILE_SIZE_BYTES:
        raise XlsxImportError("The XLSX file exceeds the 5 MB limit.")

    try:
        return load_workbook(
            io.BytesIO(file_bytes),
            read_only=True,
            data_only=True,
        )
    except (BadZipFile, InvalidFileException, OSError, ValueError) as error:
        raise XlsxImportError(
            "The XLSX file is invalid or unreadable."
        ) from error


def get_xlsx_sheet_names(file_bytes: bytes) -> tuple[str, ...]:
    workbook = _load_workbook(file_bytes)

    try:
        return tuple(sheet.title for sheet in workbook.worksheets)
    finally:
        workbook.close()

def parse_xlsx(
    file_bytes: bytes,
    sheet_name: str,
    *,
    has_header: bool = True,
) -> ParsedImportData:
    workbook = _load_workbook(file_bytes)

    try:
        worksheet = next(
            (
                sheet
                for sheet in workbook.worksheets
                if sheet.title == sheet_name
            ),
            None,
        )

        if worksheet is None:
            raise XlsxImportError("The selected worksheet is invalid.")

        rows_iterator = worksheet.iter_rows(values_only=True)

        try:
            first_row = next(rows_iterator)
        except StopIteration as error:
            raise XlsxImportError(
                "The selected worksheet has no rows."
            ) from error

        rows: list[tuple[str, ...]] = []

        if has_header:
            raw_headers = first_row
        else:
            raw_headers = ()
            rows.append(tuple(_cell_to_text(value) for value in first_row))

        for row in rows_iterator:
            if len(rows) >= MAX_DATA_ROWS:
                raise XlsxImportError(
                    "The XLSX worksheet exceeds the 10,000 row limit."
                )

            rows.append(tuple(_cell_to_text(value) for value in row))

        if has_header:
            headers = tuple(
                _cell_to_text(value).strip()
                for value in raw_headers
            )
        else:
            column_count = max((len(row) for row in rows), default=0)
            headers = tuple(
                f"Column {index + 1}"
                for index in range(column_count)
            )

        if not headers:
            raise XlsxImportError(
                "The selected worksheet has no columns."
            )

        return ParsedImportData(
            headers=headers,
            rows=tuple(rows),
        )
    finally:
        workbook.close()


def _cell_to_text(value: object) -> str:
    if value is None:
        return ""

    return str(value)