from io import BytesIO
import unittest

from openpyxl import load_workbook

from app.services.xlsx_export import build_xlsx_export


class XlsxExportTests(unittest.TestCase):
    def test_exports_headers_and_preserves_unicode_and_special_characters(self):
        file_bytes = build_xlsx_export(
            [("Merhaba, dünya", "Satır 1\nSatır 2")]
        )

        workbook = load_workbook(BytesIO(file_bytes), data_only=False)
        worksheet = workbook.active

        self.assertEqual(worksheet.title, "Cards")
        self.assertEqual(worksheet["A1"].value, "Front")
        self.assertEqual(worksheet["B1"].value, "Back")
        self.assertEqual(worksheet["A2"].value, "Merhaba, dünya")
        self.assertEqual(worksheet["B2"].value, "Satır 1\nSatır 2")

    def test_exports_formula_like_values_as_text(self):
        file_bytes = build_xlsx_export(
            [("=1+1", "+SUM(A1:A2)")]
        )

        worksheet = load_workbook(
            BytesIO(file_bytes),
            data_only=False,
        ).active

        self.assertEqual(worksheet["A2"].value, "=1+1")
        self.assertEqual(worksheet["B2"].value, "+SUM(A1:A2)")
        self.assertEqual(worksheet["A2"].data_type, "s")
        self.assertEqual(worksheet["B2"].data_type, "s")

    def test_exports_headers_when_deck_has_no_cards(self):
        file_bytes = build_xlsx_export([])

        worksheet = load_workbook(
            BytesIO(file_bytes),
            data_only=False,
        ).active

        self.assertEqual(worksheet.max_row, 1)
        self.assertEqual(
            [cell.value for cell in worksheet[1]],
            ["Front", "Back"],
        )


if __name__ == "__main__":
    unittest.main()